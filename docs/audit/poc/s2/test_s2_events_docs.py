"""S2 PoCs (events, ПГАС, documents, exports). Every test asserts the behaviour the audit found, so a PASS
means "reproduced"; if the code is fixed the test fails and should be inverted.

Run from the repo root:
    cd server && python -m pytest ../docs/audit/poc/s2/test_s2_events_docs.py -q -p no:cacheprovider
"""
import io
import os
import sys
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "..", "server", "tests"))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "..", "server"))

from conftest import (CSRF, app, clean_state, client, db, fake_eios, login_admin, login_student)  # noqa: E402,F401
import app.models as models  # noqa: E402
from app.db.database import SessionLocal  # noqa: E402
from app.services import eios, timetable  # noqa: E402

PDF = b"%PDF-1.4\n% poc\n"


@pytest.fixture(autouse=True)
def empty_catalog(app):
    s = SessionLocal()
    s.query(models.Association).delete()
    s.commit()
    s.close()


def _student(app, fake_eios, n, full_name, group="24-ИСбо-1"):
    return login_student(app, fake_eios, username=f"24-isbo-{n:03d}", eios_id=str(4000 + n), full_name=full_name, group=group)


def _uid(db, n):
    return db.query(models.User).filter_by(username=f"24-isbo-{n:03d}").one().id


def _iso(dt):
    return dt.isoformat()


def _when(days, hours=2):
    s = datetime.now(timezone.utc) + timedelta(days=days)
    return {"starts_at": _iso(s), "ends_at": _iso(s + timedelta(hours=hours))}


def _past(days_ago, hours=2):
    s = datetime.now(timezone.utc) - timedelta(days=days_ago)
    return {"starts_at": _iso(s), "ends_at": _iso(s + timedelta(hours=hours))}


@pytest.fixture
def club(app, fake_eios, db):
    """(admin, association id, leader, member, outsider)"""
    admin = login_admin(app)
    aid = admin.post("/api/v1/admin/associations", json={"name": "Медиацентр"}, headers=CSRF).json()["id"]
    leader = _student(app, fake_eios, 1, "Смирнов Макар Олегович")
    admin.put(f"/api/v1/admin/associations/{aid}/leaders/{_uid(db, 1)}", headers=CSRF)
    member = _student(app, fake_eios, 2, "Петрова Анна Сергеевна")
    member.post(f"/api/v1/associations/{aid}/apply", json={}, headers=CSRF)
    leader.post(f"/api/v1/associations/{aid}/members/{_uid(db, 2)}/decision", json={"approve": True}, headers=CSRF)
    outsider = _student(app, fake_eios, 3, "Ли Михаил Юрьевич", group="23-ИБбо-2")
    return admin, aid, leader, member, outsider


def _event(c, **data):
    r = c.post("/api/v1/events", json={"title": "Квиз", **_when(3), **data}, headers=CSRF)
    assert r.status_code == 201, r.text
    return r.json()


def _try(fn):
    """(status, exception) - the TestClient re-raises server errors, which is how a 500 shows up here."""
    try:
        return fn().status_code, None
    except Exception as exc:  # noqa: BLE001
        return None, exc


# ---- S2-001: a control character in a title makes every export that contains it fail (500) ---------

def test_poc_control_char_in_title_breaks_exports(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid, title="Квиз\x01")  # accepted: str.split() does not strip \x01
    assert ev["title"] == "Квиз\x01"
    assert leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF).status_code == 200
    # the leader's own export
    status, exc = _try(lambda: leader.get(f"/api/v1/events/{ev['id']}/export"))
    assert status is None and exc is not None, "export should have crashed"
    print("leader export:", type(exc).__name__, str(exc)[:80])
    # and the administrators' report for the whole period (one bad title takes down the report for everyone)
    start = (datetime.now(timezone.utc) - timedelta(days=1)).date().isoformat()
    end = (datetime.now(timezone.utc) + timedelta(days=30)).date().isoformat()
    for fmt in ("xlsx", "docx"):
        # needs at least one registration row to reach the writer
        status, exc = _try(lambda: admin.get("/api/v1/events/export", params={"start": start, "end": end, "format": fmt}))
        print(f"admin report {fmt}:", status, type(exc).__name__ if exc else "")
    assert _try(lambda: admin.get("/api/v1/events/export", params={"start": start, "end": end, "format": "xlsx"}))[1] is not None


def test_poc_control_char_in_manual_entry_breaks_own_export(club):
    admin, aid, leader, member, outsider = club
    r = member.post("/api/v1/achievements", json={"title": "Хакатон\x0b", "day": "2026-09-20"}, headers=CSRF)
    # \x0b is whitespace for str.split(), so use a char it keeps
    r = member.post("/api/v1/achievements", json={"title": "Хакатон\x01", "day": "2026-09-20"}, headers=CSRF)
    assert r.status_code == 201
    for fmt in ("xlsx", "docx"):
        status, exc = _try(lambda: member.get("/api/v1/portfolio/export", params={"format": fmt, "start": "2026-09-01", "end": "2027-01-31"}))
        print(f"portfolio {fmt}:", status, type(exc).__name__ if exc else "")
    assert _try(lambda: member.get("/api/v1/portfolio/export", params={"format": "xlsx", "start": "2026-09-01", "end": "2027-01-31"}))[1] is not None


# ---- formula guard works (negative result: documented as "ок") ---------------------------------------

@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r", " =", "\n="])
def test_formula_prefixes_never_become_formulas(club, db, prefix):
    admin, aid, leader, member, outsider = club
    title = f"{prefix}HYPERLINK(\"http://x\")"
    ev = _event(admin, scope="institute", title=title, **_past(2))
    admin.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    x = admin.get(f"/api/v1/events/{ev['id']}/export")
    ws = load_workbook(io.BytesIO(x.content)).active
    types = {c.data_type for row in ws.iter_rows() for c in row if c.value is not None}
    assert "f" not in types, types


# ---- S2-002: "not marked" counts as present in the ПГАС summary -----------------------------------------

def test_poc_self_registered_but_never_marked_appears_in_pgas(club):
    admin, aid, leader, member, outsider = club
    ev = _event(admin, scope="institute", title="Конференция", **_when(1))
    assert outsider.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF).status_code == 200
    # time passes: the event is over, nobody marked attendance (here: the admin moves it into the past)
    admin.put(f"/api/v1/events/{ev['id']}", json={"title": "Конференция", "scope": "institute", **_past(3)}, headers=CSRF)
    d = outsider.get(f"/api/v1/events/{ev['id']}").json()
    assert d["my_attended"] is None
    rows = outsider.get("/api/v1/portfolio").json()["rows"]
    assert [r["title"] for r in rows] == ["Конференция"], rows  # in the summary with no confirmation
    x = outsider.get("/api/v1/portfolio/export", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(x.content)).active
    assert any(c.value == "Конференция" for row in ws.iter_rows() for c in row)


# ---- S2-003: the 7-day correction window is not applied everywhere ---------------------------------------

def test_poc_fix_window_not_enforced_on_remove_role_delete_files(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(admin, scope="association", association_id=aid, title="Старое", volunteer_limit=3, **_past(40))
    admin.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    admin.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    # leader: list and attendance are frozen after a week ...
    assert leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 1)]}, headers=CSRF).status_code == 400
    assert leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": []}, headers=CSRF).status_code == 400
    # ... but these still work on a 40-day-old event:
    assert leader.patch(f"/api/v1/events/{ev['id']}/registrations/{_uid(db, 2)}", json={"role": "volunteer"}, headers=CSRF).status_code == 200
    r = leader.post(f"/api/v1/events/{ev['id']}/files", params={"name": "Позже.pdf"}, content=PDF, headers=CSRF)
    assert r.status_code == 201
    assert leader.delete(f"/api/v1/events/{ev['id']}/registrations/{_uid(db, 2)}", headers=CSRF).status_code == 200  # removes an attended person
    assert leader.delete(f"/api/v1/events/{ev['id']}", headers=CSRF).status_code == 200  # whole 40-day-old event, gone


# ---- S2-004: institute event with an organizer association: its leaders mark attendance, no bits cap ------

def test_poc_institute_event_with_organizer_is_uncapped_for_leaders(club, db):
    admin, aid, leader, member, outsider = club
    uid = _uid(db, 2)
    for i in range(12):
        ev = _event(admin, scope="institute", association_id=aid, title=f"Инст {i}", **_past(1 + 0 * i, 1))
        assert leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [uid]}, headers=CSRF).status_code == 200
        assert leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [uid]}, headers=CSRF).status_code == 200
    inst_points = member.get("/api/v1/progress").json()["facts"]["events"]
    assert inst_points == 12  # 12 x 10 bits, the 8-per-semester association cap does not apply
    # the same thing with association-level events is cut at 8
    for i in range(12):
        ev = _event(admin, scope="association", association_id=aid, title=f"Объед {i}", **_past(1, 1))
        leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [uid]}, headers=CSRF)
        leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [uid]}, headers=CSRF)
    p = member.get("/api/v1/progress").json()
    print("facts", p["facts"], "points", p["points"], "caps", p["caps"])


# ---- S2-005: organizers are not bound by the participant / volunteer limits ------------------------------

def test_poc_organizers_exceed_limits(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid, participant_limit=1, volunteer_limit=1)
    r = leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 1), _uid(db, 2)]}, headers=CSRF)
    assert r.status_code == 200 and r.json()["participants"] == 2 > r.json()["participant_limit"]
    r = leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 1), _uid(db, 2)], "role": "volunteer"}, headers=CSRF)
    assert r.json()["volunteers"] == 2 > r.json()["volunteer_limit"]


# ---- JSON exposure: who sees what in the event payload ---------------------------------------------------

def test_event_json_feedback_is_anonymous_but_contacts_reach_leaders(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid, **_past(0.3, 1))
    leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    member.post(f"/api/v1/events/{ev['id']}/feedback", json={"rating": 2, "text": "скучно"}, headers=CSRF)
    raw = leader.get(f"/api/v1/events/{ev['id']}").json()
    assert raw["feedback"] == {"count": 1, "average": 2.0, "comments": [{"rating": 2, "text": "скучно"}]}
    # an outsider put on the list by the administration shows their contacts to the association's leader
    s = SessionLocal()
    s.query(models.User).filter_by(id=_uid(db, 3)).update({"vk_url": "https://vk.com/outsider", "max_contact": "+7 000"})
    s.commit(); s.close()
    admin.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 3)]}, headers=CSRF)
    regs = leader.get(f"/api/v1/events/{ev['id']}").json()["registrations"]
    out = next(r for r in regs if r["user_id"] == _uid(db, 3))
    assert out["vk_url"] == "https://vk.com/outsider" and out["max_contact"] == "+7 000"
    # a guest sees the organizer's full name on an institute event
    inst = _event(admin, scope="institute", title="Для всех", **_when(2))
    guest = TestClient(app_of(leader)).get(f"/api/v1/events/{inst['id']}").json()
    assert guest["created_by"]


def app_of(c):
    return c.app


# ---- files: headers, Range, who can pull an order -------------------------------------------------------

def test_event_file_download_headers_range_and_self_registered_access(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(admin, scope="institute", title="Выезд", **_when(2))
    fid = admin.post(f"/api/v1/events/{ev['id']}/files", params={"name": "Распоряжение.pdf"}, content=PDF + b"x" * 200, headers=CSRF).json()["id"]
    # anyone who signs up before the start can read the order that names other students
    assert outsider.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF).status_code == 200
    r = outsider.get(f"/api/v1/attachments/{fid}")
    assert r.status_code == 200
    print({k: v for k, v in r.headers.items() if k in ("cache-control", "content-disposition", "x-content-type-options", "accept-ranges", "etag", "content-type")})
    assert r.headers["cache-control"] == "private, no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["content-disposition"].startswith("attachment")
    ranged = outsider.get(f"/api/v1/attachments/{fid}", headers={"Range": "bytes=0-3"})
    assert ranged.status_code == 206 and ranged.content == b"%PDF"
    # not signed in: no Range trick
    assert TestClient(outsider.app).get(f"/api/v1/attachments/{fid}", headers={"Range": "bytes=0-3"}).status_code == 401
    # a student who is not on the list gets 404 (and then so does a removed one)
    stranger = member  # association member? no: institute event, not registered
    assert stranger.get(f"/api/v1/attachments/{fid}").status_code == 404
    # path / extension tricks on upload
    for name in ("../../x.pdf", "a.pdf\x00.exe", "x.svg", "x.html", "x.pdf.exe"):
        r = admin.post(f"/api/v1/events/{ev['id']}/files", params={"name": name}, content=PDF, headers=CSRF)
        print(repr(name), r.status_code)


# ---- documents: /pairs asks EIOS for any year; notes can carry any name ---------------------------------

def test_poc_documents_pairs_hits_eios_for_every_year(app, fake_eios, monkeypatch):
    calls = []

    async def fetch_json(endpoint, params, timeout=5.0):
        calls.append((endpoint, tuple(sorted(params.items()))))
        return {"state": 1, "data": []}

    monkeypatch.setattr(eios, "fetch_json", fetch_json)
    timetable.clear_cache()
    c = login_student(app, fake_eios, full_name="Иванова Анна Сергеевна", group="24-ИСбо-1")
    for year in range(1990, 2030):
        assert c.get("/api/v1/documents/pairs", params={"date": f"{year}-10-01"}).status_code == 200
    print("EIOS calls from one student in 40 requests:", len(calls))
    assert len(calls) >= 40
    assert len({k for k in timetable._cache}) >= 40


def test_documents_can_be_written_for_any_name(app, fake_eios):
    c = login_student(app, fake_eios, full_name="Иванова Анна Сергеевна", group="24-ИСбо-1")
    body = {"full_name": "Петров Пётр Петрович", "group": "21-ИСбо-1", "discipline": "Математика", "control": "экзамен", "reason": "болезнь"}
    r = c.post("/api/v1/documents/retake", json=body, headers=CSRF)
    assert r.status_code == 200
    assert "Петров" in "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)


# ---- S2-006: a leader can learn who wrote an "anonymous" rating by switching attendance ----------------

def test_poc_leader_deanonymizes_rating_via_attendance(club, db, app, fake_eios):
    admin, aid, leader, member, outsider = club
    m2 = _student(app, fake_eios, 4, "Орлова Вера Игоревна")
    m2.post(f"/api/v1/associations/{aid}/apply", json={}, headers=CSRF)
    leader.post(f"/api/v1/associations/{aid}/members/{_uid(db, 4)}/decision", json={"approve": True}, headers=CSRF)
    ev = _event(leader, association_id=aid, **_past(0.3, 1))
    ids = [_uid(db, 2), _uid(db, 4)]
    leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": ids}, headers=CSRF)
    # Step 1: only Петрова is marked present; Орлова is marked absent and is locked out of rating
    leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [ids[0]]}, headers=CSRF)
    assert m2.post(f"/api/v1/events/{ev['id']}/feedback", json={"rating": 1, "text": "x"}, headers=CSRF).status_code == 400
    before = leader.get(f"/api/v1/events/{ev['id']}").json()["feedback"]["count"]
    # Step 2: Петрова rates; the only way the count can move is her
    member.post(f"/api/v1/events/{ev['id']}/feedback", json={"rating": 2, "text": "скучно"}, headers=CSRF)
    after = leader.get(f"/api/v1/events/{ev['id']}").json()["feedback"]
    assert (before, after["count"]) == (0, 1)
    assert after["comments"] == [{"rating": 2, "text": "скучно"}]  # leader now knows: Петрова gave 2, "скучно"


# ---- S2-007: /events and /calendar load every registration of every event on each call -----------------

def test_perf_event_list_scales_with_all_history(club, db):
    import time
    admin, aid, leader, member, outsider = club
    s = SessionLocal()
    now = datetime.now(timezone.utc)
    users = [models.User(username=f"perf-{i}", full_name=f"Перф {i}", hashed_password="x", group_number="24-ИСбо-1") for i in range(300)]
    s.add_all(users)
    s.commit()
    uids = [u.id for u in users]
    for e in range(200):  # ~ two years of events, each with 150 people (a whole group or two)
        ev = models.Event(scope="institute", title=f"Событие {e}", starts_at=now - timedelta(days=e * 3), ends_at=now - timedelta(days=e * 3) + timedelta(hours=2))
        s.add(ev)
        s.flush()
        s.add_all(models.EventRegistration(event_id=ev.id, user_id=uids[(e + k) % 300], source="admin_group") for k in range(150))
    s.commit()
    s.close()
    t = time.perf_counter()
    r = outsider.get("/api/v1/events", params={"view": "upcoming"})
    dt = time.perf_counter() - t
    t = time.perf_counter()
    c = outsider.get("/api/v1/calendar", params={"start": now.date().isoformat(), "end": (now + timedelta(days=7)).date().isoformat()})
    dc = time.perf_counter() - t
    print(f"\nGET /events (upcoming, nothing upcoming): {dt:.2f}s status={r.status_code} rows={len(r.json())}; /calendar {dc:.2f}s status={c.status_code}")


# ---- S2-003 (cont.): re-dating an old event reopens the 7-day window ---------------------------------

def test_poc_leader_redates_old_event_into_the_window(club, db):
    admin, aid, leader, member, outsider = club
    ev = _event(admin, scope="association", association_id=aid, title="Давнее", **_past(40))
    assert leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF).status_code == 400
    again = _past(0.3, 1)
    r = leader.put(f"/api/v1/events/{ev['id']}", json={"title": "Давнее", "scope": "association", "association_id": aid, **again}, headers=CSRF)
    assert r.status_code == 200, r.text  # a 40-day-old event now "happened" yesterday
    assert leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF).status_code == 200
    assert leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [_uid(db, 2)]}, headers=CSRF).status_code == 200


# ---- documents: markup in user text is escaped in DOCX and PDF (negative result, "ок") ------------------

def test_documents_escape_markup(app, fake_eios, monkeypatch):
    c = login_student(app, fake_eios, full_name="Иванова Анна Сергеевна", group="24-ИСбо-1")
    evil = "<a href='http://evil.example'>x</a> <b>y</b> & <font color=red>z</font> <img src='file:///etc/passwd'/> ]]> <w:p/>"
    body = {"full_name": "Иванова Анна Сергеевна", "group": "<b>24</b>&<i>", "date_from": "2026-09-23", "reason": evil, "attachment": evil}
    pdf = c.post("/api/v1/documents/explanatory", params={"format": "pdf"}, json=body, headers=CSRF)
    assert pdf.status_code == 200 and b"/URI" not in pdf.content and b"evil.example" not in pdf.content.replace(b" ", b"")
    docx = c.post("/api/v1/documents/explanatory", params={"format": "docx"}, json=body, headers=CSRF)
    assert docx.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert "<a href='http://evil.example'>x</a>" in text  # literal text, not markup
    # a control character in a free field is stripped, not a crash
    body["reason"] = "болезнь\x01\x00"
    assert c.post("/api/v1/documents/explanatory", params={"format": "docx"}, json=body, headers=CSRF).status_code == 200


# ---- BFLA / BOLA matrix on the organizer endpoints (negative result, "ок") -----------------------------

def test_organizer_endpoints_reject_everyone_else(club, db, app, fake_eios):
    admin, aid, leader, member, outsider = club
    aid2 = admin.post("/api/v1/admin/associations", json={"name": "Другое"}, headers=CSRF).json()["id"]
    other = _student(app, fake_eios, 5, "Кузнецов Олег Иванович")
    admin.put(f"/api/v1/admin/associations/{aid2}/leaders/{_uid(db, 5)}", headers=CSRF)
    ev = _event(leader, association_id=aid, volunteer_limit=2, **_when(2))
    leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    inst = _event(admin, scope="institute", title="Общее", **_when(2))
    e, i, uid = ev["id"], inst["id"], _uid(db, 2)
    calls = [
        ("put", f"/events/{e}", {"json": {"title": "x", **_when(2), "association_id": aid}}),
        ("delete", f"/events/{e}", {}),
        ("post", f"/events/{e}/registrations", {"json": {"user_ids": [uid]}}),
        ("post", f"/events/{e}/groups", {"json": {"groups": ["24-ИСбо-1"]}}),
        ("delete", f"/events/{e}/registrations/{uid}", {}),
        ("patch", f"/events/{e}/registrations/{uid}", {"json": {"role": "volunteer"}}),
        ("delete", f"/events/{e}/removals/{uid}", {}),
        ("put", f"/events/{e}/attendance", {"json": {"user_ids": []}}),
        ("get", f"/events/{e}/export", {}),
        ("post", f"/events/{e}/files?name=a.pdf", {"content": PDF}),
        ("post", f"/events/{e}/links", {"json": {"url": "https://example.org/x"}}),
        # institute event without an organizer: administrators only
        ("put", f"/events/{i}", {"json": {"title": "x", "scope": "institute", **_when(2)}}),
        ("delete", f"/events/{i}", {}),
        ("post", f"/events/{i}/registrations", {"json": {"user_ids": [uid]}}),
        ("put", f"/events/{i}/attendance", {"json": {"user_ids": []}}),
        ("get", f"/events/{i}/export", {}),
        ("get", "/events/export?start=2026-01-01&end=2026-12-31", {}),
    ]
    out = []
    for who, c in (("anon", TestClient(app)), ("outsider", outsider), ("member", member), ("foreign leader", other)):
        for method, path, kw in calls:
            r = getattr(c, method)(f"/api/v1{path}", headers=CSRF, **kw)
            out.append((who, method, path, r.status_code))
            assert r.status_code in (401, 403, 404), (who, method, path, r.status_code, r.text[:100])
    # the member of the association may read the event but nothing more
    assert member.get(f"/api/v1/events/{e}").status_code == 200
    # after the leader is dropped from the association the management rights go with it
    assert admin.delete(f"/api/v1/admin/associations/{aid}/leaders/{_uid(db, 1)}", headers=CSRF).status_code == 200
    r = leader.put(f"/api/v1/events/{e}", json={"title": "x", **_when(2), "association_id": aid}, headers=CSRF)
    assert r.status_code == 403
    assert leader.get(f"/api/v1/events/{e}/export").status_code == 403
    # and a deactivated association hides its events from everyone but administrators
    admin.put(f"/api/v1/admin/associations/{aid}", json={"name": "Медиацентр", "is_active": False}, headers=CSRF)
    print("member sees event of a deactivated association:", member.get(f"/api/v1/events/{e}").status_code)
