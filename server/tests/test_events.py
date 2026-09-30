import io
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from openpyxl import load_workbook

import app.models as models
from app.db.database import SessionLocal
from conftest import CSRF, login_admin, login_student

PDF = b"%PDF-1.4\n% test\n"


@pytest.fixture(autouse=True)
def empty_catalog(app):
    session = SessionLocal()
    session.query(models.Association).delete()
    session.commit()
    session.close()


def _student(app, fake_eios, n, full_name, group="24-ИСбо-1"):
    return login_student(app, fake_eios, username=f"24-isbo-{n:03d}", eios_id=str(4000 + n),
                         full_name=full_name, group=group)


def _uid(db, n):
    return db.query(models.User).filter_by(username=f"24-isbo-{n:03d}").one().id


def _iso(dt):
    return dt.isoformat()


def _when(days, hours=2):
    start = datetime.now(timezone.utc) + timedelta(days=days)
    return {"starts_at": _iso(start), "ends_at": _iso(start + timedelta(hours=hours))}


@pytest.fixture
def club(app, fake_eios, db):
    """(admin, association id, leader, member, outsider)"""
    admin = login_admin(app)
    aid = admin.post("/api/v1/admin/associations", json={"name": "Медиацентр"}, headers=CSRF).json()["id"]
    leader = _student(app, fake_eios, 1, "Смирнов Макар Олегович")
    admin.put(f"/api/v1/admin/associations/{aid}/leaders/{_uid(db, 1)}", headers=CSRF)
    member = _student(app, fake_eios, 2, "Петрова Анна Сергеевна")
    member.post(f"/api/v1/associations/{aid}/apply", json={}, headers=CSRF)
    leader.post(f"/api/v1/associations/{aid}/members/{_uid(db, 2)}/decision", json={"approve": True}, headers=CSRF)
    outsider = _student(app, fake_eios, 3, "Ли Михаил Юрьевич", group="23-ИБбо-2")
    return admin, aid, leader, member, outsider


def _event(client, **data):
    r = client.post("/api/v1/events", json={"title": "Квиз", **_when(3), **data}, headers=CSRF)
    assert r.status_code == 201, r.text
    return r.json()


def test_association_and_institute_events_visibility(club, app):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid, participant_limit=1)
    assert ev["can_manage"] and ev["scope"] == "association"
    assert [e["title"] for e in member.get("/api/v1/events").json()] == ["Квиз"]
    assert outsider.get("/api/v1/events").json() == []
    assert outsider.get(f"/api/v1/events/{ev['id']}").status_code == 404
    # Institute events: administrators only
    r = leader.post("/api/v1/events", json={"title": "День ИВИТШ", "scope": "institute", **_when(5)}, headers=CSRF)
    assert r.status_code == 403
    inst = _event(admin, title="День ИВИТШ", scope="institute", volunteer_limit=2)
    assert inst["association"] is None
    assert {e["title"] for e in outsider.get("/api/v1/events").json()} == {"День ИВИТШ"}
    # Guests see the institute's events too, without anyone's registrations
    from fastapi.testclient import TestClient
    guest = TestClient(app).get(f"/api/v1/events/{inst['id']}").json()
    assert guest["title"] == "День ИВИТШ" and guest["my_role"] is None
    # A leader cannot turn their event into an institute one
    r = leader.put(f"/api/v1/events/{ev['id']}", json={"title": "Квиз", "scope": "institute", "association_id": aid, **_when(3)}, headers=CSRF)
    assert r.status_code == 403


def test_registration_limits_and_roles(club):
    admin, aid, leader, member, outsider = club
    ev = _event(admin, title="День ИВИТШ", scope="institute", participant_limit=1, volunteer_limit=1)
    r = outsider.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF)
    assert r.json()["my_role"] == "participant" and r.json()["participants"] == 1
    assert member.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF).status_code == 409
    r = member.post(f"/api/v1/events/{ev['id']}/register", json={"role": "volunteer"}, headers=CSRF)
    assert r.json()["my_role"] == "volunteer"
    # Switching roles frees the old place
    assert outsider.post(f"/api/v1/events/{ev['id']}/register", json={"role": "volunteer"}, headers=CSRF).status_code == 409
    assert member.delete(f"/api/v1/events/{ev['id']}/register", headers=CSRF).json()["my_role"] is None
    no_volunteers = _event(leader, association_id=aid)
    assert member.post(f"/api/v1/events/{no_volunteers['id']}/register", json={"role": "volunteer"}, headers=CSRF).status_code == 400
    closed = _event(leader, association_id=aid, registration_open=False)
    assert member.post(f"/api/v1/events/{closed['id']}/register", json={}, headers=CSRF).status_code == 400


def test_leaders_add_members_only_and_groups_are_for_admins(club, db, app, fake_eios):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid)
    r = leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 3)]}, headers=CSRF)
    assert r.status_code == 400
    r = leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    assert [(p["full_name"], p["source"]) for p in r.json()["registrations"]] == [("Петрова Анна Сергеевна", "leader")]
    # Only the administration enrolls whole groups
    body = {"groups": ["24-исбо-1", "99-НЕТ-1"]}
    assert leader.post(f"/api/v1/events/{ev['id']}/groups", json=body, headers=CSRF).status_code == 403
    _student(app, fake_eios, 4, "Орлова Вера Игоревна")
    r = admin.post(f"/api/v1/events/{ev['id']}/groups", json=body, headers=CSRF).json()
    assert r == {"added": 2, "already": 1, "unknown_groups": ["99-НЕТ-1"]}  # the leader and Орлова; Петрова was there
    regs = admin.get(f"/api/v1/events/{ev['id']}").json()["registrations"]
    assert sorted(p["source"] for p in regs) == ["admin_group", "admin_group", "leader"]


def test_attendance_feedback_and_documents(club, db):
    admin, aid, leader, member, outsider = club
    past = datetime.now(timezone.utc) - timedelta(days=1)
    ev = _event(leader, association_id=aid, starts_at=_iso(past), ends_at=_iso(past + timedelta(hours=2)))
    assert member.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF).status_code == 400  # already over
    leader.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2), _uid(db, 1)]}, headers=CSRF)
    r = leader.put(f"/api/v1/events/{ev['id']}/attendance", json={"user_ids": [_uid(db, 2)]}, headers=CSRF)
    assert {p["full_name"]: p["attended"] for p in r.json()["registrations"]} == {
        "Петрова Анна Сергеевна": True, "Смирнов Макар Олегович": False}

    assert [e["id"] for e in member.get("/api/v1/events/feedback/pending").json()] == [ev["id"]]
    assert leader.get("/api/v1/events/feedback/pending").json() == []  # marked absent
    r = member.post(f"/api/v1/events/{ev['id']}/feedback", json={"rating": 5, "text": "Круто"}, headers=CSRF)
    assert r.json()["my_feedback"] == {"rating": 5, "text": "Круто"}
    assert member.post(f"/api/v1/events/{ev['id']}/feedback", json={"rating": 6}, headers=CSRF).status_code == 422
    assert member.get("/api/v1/events/feedback/pending").json() == []
    summary = leader.get(f"/api/v1/events/{ev['id']}").json()["feedback"]
    assert summary == {"count": 1, "average": 5.0, "comments": [{"rating": 5, "text": "Круто"}]}  # anonymous

    r = leader.post(f"/api/v1/events/{ev['id']}/files", params={"name": "Распоряжение.pdf"}, content=PDF, headers=CSRF)
    assert r.status_code == 201, r.text
    fid = r.json()["id"]
    assert member.post(f"/api/v1/events/{ev['id']}/files", params={"name": "x.pdf"}, content=PDF, headers=CSRF).status_code == 403
    assert member.get(f"/api/v1/attachments/{fid}").content == PDF
    assert [a["title"] for a in member.get(f"/api/v1/events/{ev['id']}").json()["attachments"]] == ["Распоряжение.pdf"]
    assert outsider.get(f"/api/v1/attachments/{fid}").status_code == 404

    x = leader.get(f"/api/v1/events/{ev['id']}/export")
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(x.content)).active
    names = [row[1] for row in sheet.iter_rows(min_row=5, values_only=True)]
    assert names == ["Петрова Анна Сергеевна", "Смирнов Макар Олегович"]
    assert member.get(f"/api/v1/events/{ev['id']}/export").status_code == 403


def test_pgas_summary_with_manual_entries_and_exports(club, db):
    admin, aid, leader, member, outsider = club
    day = datetime(2026, 9, 20, 12, tzinfo=timezone.utc)
    ev = _event(admin, title="=Хакатон", scope="institute", starts_at=_iso(day), ends_at=_iso(day + timedelta(hours=5)),
                volunteer_limit=3)
    admin.post(f"/api/v1/events/{ev['id']}/registrations", json={"user_ids": [_uid(db, 2)], "role": "volunteer"}, headers=CSRF)
    admin.post(f"/api/v1/events/{ev['id']}/files", params={"name": "Благодарность.pdf"}, content=PDF, headers=CSRF)
    r = member.post("/api/v1/achievements", json={"title": "Олимпиада по программированию", "organizer": "КГУ",
                                                  "day": "2026-10-10", "role": "Призёр (2 место)"}, headers=CSRF)
    assert r.status_code == 201, r.text
    ach = r.json()
    assert member.post(f"/api/v1/achievements/{ach['id']}/files", params={"name": "скан.pdf"}, content=PDF, headers=CSRF).status_code == 201
    assert outsider.put(f"/api/v1/achievements/{ach['id']}", json={"title": "x", "day": "2026-10-10"}, headers=CSRF).status_code == 404

    p = member.get("/api/v1/portfolio", params={"start": "2026-09-01", "end": "2027-01-31"}).json()
    assert [(r["kind"], r["title"], r["role"], r["level"]) for r in p["rows"]] == [
        ("event", "=Хакатон", "Волонтёр", "Институт"),
        ("manual", "Олимпиада по программированию", "Призёр (2 место)", "Внесено студентом"),
    ]
    assert [d["title"] for d in p["rows"][0]["documents"]] == ["Благодарность.pdf"]
    assert outsider.get("/api/v1/portfolio", params={"start": "2026-09-01", "end": "2027-01-31"}).json()["rows"] == []

    x = member.get("/api/v1/portfolio/export", params={"format": "xlsx", "start": "2026-09-01", "end": "2027-01-31"})
    sheet = load_workbook(io.BytesIO(x.content)).active
    assert sheet["A3"].value == "За осенний семестр 2026/2027 учебного года"
    assert sheet["C6"].value == "'=Хакатон"  # never a formula
    d = member.get("/api/v1/portfolio/export", params={"format": "docx", "start": "2026-09-01", "end": "2027-01-31"})
    doc = Document(io.BytesIO(d.content))
    assert "Петрова Анна Сергеевна" in doc.paragraphs[1].text
    assert doc.tables[0].rows[2].cells[2].text == "Олимпиада по программированию"

    # Deleting the entry removes its scan from disk too
    assert member.delete(f"/api/v1/achievements/{ach['id']}", headers=CSRF).status_code == 200
    assert db.query(models.Attachment).filter(models.Attachment.achievement_id.isnot(None)).count() == 0


def test_events_show_up_in_the_calendar(club):
    admin, aid, leader, member, outsider = club
    ev = _event(leader, association_id=aid, **_when(1))
    start = (datetime.now(timezone.utc) + timedelta(days=1)).date().isoformat()
    member.post(f"/api/v1/events/{ev['id']}/register", json={}, headers=CSRF)
    items = member.get("/api/v1/calendar", params={"start": start, "days": 2}).json()["items"]
    assert [(i["type"], i["title"], i["status"]) for i in items if i["type"] == "event"] == [("event", "Квиз", "participant")]
    assert [i for i in outsider.get("/api/v1/calendar", params={"start": start, "days": 2}).json()["items"] if i["type"] == "event"] == []


def test_only_the_main_admin_manages_administrators(app, fake_eios, db):
    main = login_admin(app)
    _student(app, fake_eios, 7, "Егоров Егор Егорович")
    _student(app, fake_eios, 8, "Тимофеева Вера Павловна")
    assert main.patch(f"/api/v1/admin/users/{_uid(db, 7)}/role", json={"role": "admin"}, headers=CSRF).status_code == 200
    egor = _student(app, fake_eios, 7, "Егоров Егор Егорович")
    # A second administrator manages students but not administrators
    assert egor.patch(f"/api/v1/admin/users/{_uid(db, 8)}/role", json={"role": "moderator"}, headers=CSRF).status_code == 200
    assert egor.patch(f"/api/v1/admin/users/{_uid(db, 8)}/role", json={"role": "admin"}, headers=CSRF).status_code == 403
    main.patch(f"/api/v1/admin/users/{_uid(db, 8)}/role", json={"role": "admin"}, headers=CSRF)
    assert egor.patch(f"/api/v1/admin/users/{_uid(db, 8)}/role", json={"role": "student"}, headers=CSRF).status_code == 403
    assert egor.patch(f"/api/v1/admin/users/{_uid(db, 8)}/block", json={"blocked": True}, headers=CSRF).status_code == 403
    assert egor.delete(f"/api/v1/admin/users/{_uid(db, 8)}", headers=CSRF).status_code == 403
    assert main.patch(f"/api/v1/admin/users/{_uid(db, 8)}/role", json={"role": "student"}, headers=CSRF).status_code == 200
