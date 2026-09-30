"""Semester bounds and the Excel / Word files: the student's ПГАС summary and an event's participant list."""
import io
from datetime import date
from typing import Iterable, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUMMARY_COLUMNS = ("№", "Дата", "Мероприятие", "Организатор", "Уровень", "Роль", "Подтверждающие документы")


def semester_of(day: date) -> Tuple[date, date]:
    """Autumn: 1 September – 31 January; spring: 1 February – 31 August."""
    if day.month >= 9:
        return date(day.year, 9, 1), date(day.year + 1, 1, 31)
    if day.month == 1:
        return date(day.year - 1, 9, 1), date(day.year, 1, 31)
    return date(day.year, 2, 1), date(day.year, 8, 31)


def semester_title(start: date, end: date) -> str:
    if (start.month, start.day, end.month, end.day) == (9, 1, 1, 31):
        return f"осенний семестр {start.year}/{end.year} учебного года"
    if (start.month, start.day, end.month, end.day) == (2, 1, 8, 31):
        return f"весенний семестр {start.year - 1}/{start.year} учебного года"
    return f"период с {start:%d.%m.%Y} по {end:%d.%m.%Y}"


# Characters Excel would read as a formula when a cell starts with them
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def _sheet(title: str, header: Sequence[str], rows: Iterable[Sequence], widths: Sequence[int], heading: List[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for line in heading:
        ws.append([_cell(line)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=ws.max_row == 1, size=13 if ws.max_row == 1 else 11)
    if heading:
        ws.append([])
    ws.append(list(header))
    head_row = ws.max_row
    for col, _ in enumerate(header, start=1):
        c = ws.cell(row=head_row, column=col)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E4ECF9")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_cell(v) for v in row])
        for col in range(1, len(header) + 1):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary_rows(rows) -> List[Tuple]:
    return [
        (n, r["day"].strftime("%d.%m.%Y"), r["title"], r["organizer"], r["level"], r["role"],
         ", ".join(d["title"] for d in r["documents"]) or "—")
        for n, r in enumerate(rows, start=1)
    ]


def summary_xlsx(full_name: str, group: Optional[str], start: date, end: date, rows) -> bytes:
    heading = [
        "Участие в мероприятиях",
        f"{full_name}{', группа ' + group if group else ''}",
        f"За {semester_title(start, end)}",
    ]
    return _sheet("ПГАС", SUMMARY_COLUMNS, summary_rows(rows), (5, 12, 40, 28, 22, 16, 36), heading)


def summary_docx(full_name: str, group: Optional[str], start: date, end: date, rows, today: date) -> bytes:
    doc = Document()
    section = doc.sections[0]
    # Landscape: seven columns
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.8))
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Сведения об участии в мероприятиях")
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph(f"Студент: {full_name}" + (f", группа {group}" if group else ""))
    doc.add_paragraph(f"Период: {semester_title(start, end)}")

    table = doc.add_table(rows=1, cols=len(SUMMARY_COLUMNS))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, SUMMARY_COLUMNS):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in summary_rows(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    if not rows:
        doc.add_paragraph("За этот период мероприятий нет.")
    doc.add_paragraph()
    doc.add_paragraph(f"Сформировано на портале «ИВИТШ Хаб» {today:%d.%m.%Y}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ROLE_TEXT = {"participant": "Участник", "volunteer": "Волонтёр"}
SOURCE_TEXT = {"self": "Сам(а)", "leader": "Руководитель", "admin": "Администрация", "admin_group": "Группой (администрация)"}


def participants_xlsx(title: str, when: str, registrations) -> bytes:
    rows = [
        (n, r.user.full_name, r.user.group_number or "", ROLE_TEXT.get(r.role, r.role), SOURCE_TEXT.get(r.source, r.source),
         {True: "Да", False: "Нет"}.get(r.attended, "Не отмечено"))
        for n, r in enumerate(registrations, start=1)
    ]
    return _sheet("Участники", ("№", "ФИО", "Группа", "Роль", "Кто записал", "Был(а)"), rows,
                  (5, 36, 14, 14, 24, 14), [title, when])
